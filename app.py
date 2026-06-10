from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import io
import requests as req_lib

app = Flask(__name__)
CORS(app, origins='*')

@app.after_request
def after_request(response):
    response.headers.add('Access-Control-Allow-Origin', '*')
    response.headers.add('Access-Control-Allow-Headers', 'Content-Type')
    response.headers.add('Access-Control-Allow-Methods', 'GET,POST,OPTIONS')
    return response

OR  = RGBColor(0xFF, 0x5B, 0x00)
DK  = RGBColor(0x21, 0x27, 0x2A)
WH  = RGBColor(0xFF, 0xFF, 0xFF)
GR  = RGBColor(0xF0, 0xF0, 0xF0)
BK  = RGBColor(0x00, 0x00, 0x00)
GY2 = RGBColor(0xEA, 0xE8, 0xE9)
DK2 = RGBColor(0x2A, 0x2A, 0x2A)
LGY = RGBColor(0x55, 0x55, 0x55)
MGY = RGBColor(0xBB, 0xBB, 0xBB)
OR2 = RGBColor(0xFB, 0x79, 0x01)

LOGO_W = 'https://res.cloudinary.com/dpcojkrta/image/upload/v1780945082/Logo_Kruger_naranja_blanco_negro_y_plomo-01_mmdzkk.png'
LOGO_B = 'https://res.cloudinary.com/dpcojkrta/image/upload/v1780944928/Logo_Kruger_naranja_blanco_negro_y_plomo-02_ynrqui.png'

# Imágenes fijas por slide (orden definido)
SLIDE_IMGS = {
    'portada':      'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100423/Imagen1_hsptrs.png',
    'disclaimer':   'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100456/Imagen2_a51brf.png',
    'quienes':      'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100548/Imagen4_phr6cl.png',
    'porque':       'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100592/Imagen5_nxwp0o.png',
    'como':         'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100646/Imagen8_cdxxs3.png',
    'dolores':      'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100690/Imagen14_uzhsye.png',
    'case_intro_a': 'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100753/Imagen9_kechhv.png',
    'case_detail_a':'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100802/Imagen10_lhl5is.png',
    'case_intro_b': 'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100834/Imagen11_qjevxo.png',
    'case_intro_b2':'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100881/Imagen12_jj8nsf.png',
    'case_detail_b':'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100894/Imagen13_bddubq.png',
    'gancho_main':  'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100894/Imagen13_bddubq.png',
    'gancho_side':  'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100952/Imagen16_xl4ujj.png',
    'cierre':       'https://res.cloudinary.com/dpcojkrta/image/upload/v1781100952/Imagen16_xl4ujj.png',
}

# Partner images en orden
PARTNER_IMGS = [
    'https://res.cloudinary.com/dpcojkrta/image/upload/v1781019937/Imagen8_ljshve.png',
    'https://res.cloudinary.com/dpcojkrta/image/upload/v1781019958/Imagen7_vlvu1d.png',
    'https://res.cloudinary.com/dpcojkrta/image/upload/v1781019980/Imagen6_do5kih.png',
    'https://res.cloudinary.com/dpcojkrta/image/upload/v1781019998/Imagen5_cvmpwy.png',
    'https://res.cloudinary.com/dpcojkrta/image/upload/v1781020018/Imagen3_wae64w.png',
    'https://res.cloudinary.com/dpcojkrta/image/upload/v1781020033/Imagen4_ehxzjc.png',
    'https://res.cloudinary.com/dpcojkrta/image/upload/v1781020053/Imagen2_l0jovf.png',
    'https://res.cloudinary.com/dpcojkrta/image/upload/v1781020068/Imagen1_fdslxh.png',
]

W = Inches(13.33)
H = Inches(7.5)

def i(v): return Inches(v)
def p(v): return Pt(v)
def hc(h):
    h = h.lstrip('#')
    return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

def fetch_image(url):
    r = req_lib.get(url, timeout=10)
    return io.BytesIO(r.content)

def set_bg(sl, rgb):
    fill = sl.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb

def add_rect(sl, x, y, w, h, fill_rgb, line_rgb=None):
    shape = sl.shapes.add_shape(1, i(x), i(y), i(w), i(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if line_rgb:
        shape.line.color.rgb = line_rgb
        shape.line.width = Pt(0.5)
    else:
        shape.line.fill.background()
    return shape

def add_text(sl, x, y, w, h, text, size, bold=False, color=None,
             align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txb = sl.shapes.add_textbox(i(x), i(y), i(w), i(h))
    tf = txb.text_frame
    tf.word_wrap = wrap
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = p(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = 'Inter'
    if color: run.font.color.rgb = color
    return txb

def add_runs(sl, x, y, w, h, runs, align=PP_ALIGN.LEFT, wrap=True):
    txb = sl.shapes.add_textbox(i(x), i(y), i(w), i(h))
    tf = txb.text_frame
    tf.word_wrap = wrap
    para = tf.paragraphs[0]
    para.alignment = align
    for txt, sz, bold, clr in runs:
        run = para.add_run()
        run.text = txt
        run.font.size = p(sz)
        run.font.bold = bold
        run.font.name = 'Inter'
        if clr: run.font.color.rgb = clr
    return txb

def add_logo(sl, url, x, y, w, h):
    try:
        sl.shapes.add_picture(fetch_image(url), i(x), i(y), i(w), i(h))
    except Exception as e:
        print(f"Logo error: {e}")

def add_ph(sl, x, y, w, h, label):
    shape = sl.shapes.add_shape(1, i(x), i(y), i(w), i(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = hc('F5F5F5')
    shape.line.color.rgb = OR
    shape.line.width = Pt(1)
    add_text(sl, x, y, w, h, label, 9, bold=True,
             color=hc('888888'), align=PP_ALIGN.CENTER)

# ── TEXTOS BILINGÜE ───────────────────────────────────────────
def T(lang, es, en):
    return en if lang == 'en' else es

# ── SLIDES ────────────────────────────────────────────────────

def slide_portada(prs, sl, cfg, lang):
    set_bg(sl, OR)
    add_logo(sl, SLIDE_IMGS['portada'], 5.98, 0, 7.35, 7.5)
    add_logo(sl, LOGO_W, 0.4, 0.25, 2.5, 0.62)
    add_text(sl, 8.5, 0.22, 4.6, 0.4, cfg.get('name',''), 13, bold=True, color=WH, align=PP_ALIGN.RIGHT)
    add_text(sl, 8.5, 0.6, 4.6, 0.35, cfg.get('title',''), 11, color=hc('FFD0B0'), align=PP_ALIGN.RIGHT)
    if lang == 'en':
        add_text(sl, 0.4, 2.2, 5.4, 3.2,
                 'WEB & MOBILE APP\nDEVELOPMENT\nCASE STUDIES',
                 42, bold=True, color=WH)
    else:
        add_text(sl, 0.4, 2.2, 5.4, 3.2,
                 'DESARROLLO DE\nAPLICACIONES\nWEB & MOBILE\nCASOS DE ÉXITO',
                 38, bold=True, color=WH)
    if lang == 'en':
        add_runs(sl, 0.4, 5.9, 5.4, 0.9, [
            ('Designing scalable digital experiences for ', 17, False, WH),
            ('enterprise, telecom, fintech and customer ecosystems.', 17, True, WH),
        ], wrap=True)
    else:
        add_runs(sl, 0.4, 5.9, 5.4, 0.9, [
            ('Diseñando experiencias digitales escalables para ', 17, False, WH),
            ('enterprise, telecom, fintech y ecosistemas de cliente.', 17, True, WH),
        ], wrap=True)

def slide_disclaimer(prs, sl, lang):
    set_bg(sl, GY2)
    add_logo(sl, SLIDE_IMGS['disclaimer'], 7.8, 0, 5.53, 7.5)
    add_logo(sl, LOGO_B, 0.75, 0.5, 2.2, 0.55)
    add_text(sl, 0.75, 1.25, 6.5, 1.0, 'Disclaimer', 47, bold=True, color=OR)
    add_rect(sl, 0.75, 2.3, 0.85, 0.05, DK)
    if lang == 'en':
        add_runs(sl, 0.75, 2.5, 6.2, 3.8, [
            ('The information contained in this document is ', 17, False, DK),
            ('confidential, privileged, and only for the information of the intended recipient', 17, True, DK),
            (', and may not be used, published, or redistributed ', 17, False, DK),
            ('without the prior written consent ', 17, True, DK),
            ('of Kruger Corporation.', 17, True, DK),
        ], wrap=True)
    else:
        add_runs(sl, 0.75, 2.5, 6.2, 3.8, [
            ('La información contenida en este documento es ', 17, False, DK),
            ('confidencial, privilegiada y solo para uso del destinatario previsto', 17, True, DK),
            (', y no puede ser usada, publicada ni redistribuida ', 17, False, DK),
            ('sin el consentimiento previo por escrito ', 17, True, DK),
            ('de Kruger Corporation.', 17, True, DK),
        ], wrap=True)

def slide_quienes(prs, sl, lang):
    set_bg(sl, GR)
    add_rect(sl, 12.33, 0, 1.0, 1.0, OR)
    add_logo(sl, LOGO_B, 0.6, 0.3, 2.2, 0.55)
    add_rect(sl, 0.6, 1.0, 1.5, 0.05, OR)
    add_runs(sl, 0.6, 1.15, 6.5, 0.4, [
        (T(lang,'BUILDING SOLUTIONS. ','BUILDING SOLUTIONS. '), 14, True, DK),
        (T(lang,'DELIVERING SUCCESS.','DELIVERING SUCCESS.'), 14, True, OR),
    ])
    if lang == 'en':
        add_runs(sl, 0.6, 1.65, 5.8, 1.4, [
            ('Kruger creates web and mobile applications that help companies turn ', 19, False, DK),
            ('complex digital products', 19, True, DK),
            (' into simple, user-centered experiences.', 19, False, DK),
        ], wrap=True)
    else:
        add_runs(sl, 0.6, 1.65, 5.8, 1.4, [
            ('Kruger crea aplicaciones web y móviles que ayudan a las empresas a convertir ', 19, False, DK),
            ('productos digitales complejos', 19, True, DK),
            (' en experiencias simples y centradas en el usuario.', 19, False, DK),
        ], wrap=True)
    stats_es = [('30+','AÑOS DE INNOVACIÓN'),('1500+','PROYECTOS ENTREGADOS'),
                ('450+','EXPERTOS CERTIFICADOS'),('14','PAÍSES DE PRESENCIA GLOBAL')]
    stats_en = [('30+','YEARS OF INNOVATION'),('1500+','PROJECTS DELIVERED'),
                ('450+','CERTIFIED EXPERTS'),('14','COUNTRIES WORLDWIDE')]
    stats = stats_en if lang == 'en' else stats_es
    for idx,(v,l) in enumerate(stats):
        y = 3.2 + idx * 0.95
        add_rect(sl, 0.6, y, 0.38, 0.38, WH, OR)
        add_text(sl, 1.15, y, 5.5, 0.38, f'{v}  {l}', 14, bold=True, color=OR)
    add_logo(sl, SLIDE_IMGS['quienes'], 7.1, 0.3, 5.9, 6.9)

def slide_porque(prs, sl, lang):
    set_bg(sl, OR)
    add_text(sl, 7.2, -0.3, 5.5, 4.5, '01', 227, bold=True, color=WH)
    add_text(sl, 6.9, 5.2, 6.0, 1.1, T(lang,'WHY','WHY'), 60, color=WH)
    add_text(sl, 6.9, 6.1, 6.3, 1.3, 'KRUGER', 75, bold=True, color=WH)
    puntos_es = [
        'Transformamos negocios con tecnología centrada en personas',
        '14 países · 30 años · 1500+ proyectos entregados',
        'Expertos en sectores regulados y ecosistemas enterprise'
    ]
    puntos_en = [
        'We transform businesses with human-centered technology',
        '14 countries · 30 years · 1500+ projects delivered',
        'Experts in regulated sectors and enterprise ecosystems'
    ]
    puntos = puntos_en if lang == 'en' else puntos_es
    for idx, pt in enumerate(puntos):
        add_runs(sl, 7.2, 2.6 + idx * 0.78, 5.9, 0.72, [
            (f'{idx+1}.  ', 13, True, hc('FFE0CC')),
            (pt, 13, False, hc('FFE0CC')),
        ], wrap=True)
    add_logo(sl, LOGO_W, 11.5, 6.85, 1.6, 0.45)
    add_logo(sl, SLIDE_IMGS['porque'], 0.3, 0.3, 6.3, 6.9)

def slide_como(prs, sl, lang):
    set_bg(sl, GR)
    add_rect(sl, 12.33, 0, 1.0, 1.0, OR)
    add_runs(sl, 0.6, 0.2, 7.5, 1.3, [
        (T(lang,'What ','What '), 63, True, DK),
        (T(lang,'We Build','We Build'), 63, True, OR),
    ])
    if lang == 'en':
        add_runs(sl, 0.6, 1.6, 7.0, 0.55, [
            ('We build ', 17, False, LGY),
            ('scalable digital experiences', 17, True, DK),
            (' that connect ', 17, False, LGY),
            ('technology, operations and customer engagement.', 17, True, DK),
        ], wrap=True)
        items = [
            ('Mobile & Self-Service Apps', 'Digital journeys for customers.'),
            ('Payment & Transactional UX', 'Secure flows for conversion.'),
            ('Enterprise Portals', 'Complex processes simplified.'),
            ('B2C Digital Experiences', 'Focused on usability.'),
            ('B2B Platforms', 'Operational tools for clarity.'),
            ('UX/UI Systems', 'Multi-screen interfaces.'),
        ]
        footer = ('HUMAN-CENTERED ', 'digital products powered by INFINITE DIGITAL TECHNOLOGY.')
    else:
        add_runs(sl, 0.6, 1.6, 7.0, 0.55, [
            ('Construimos ', 17, False, LGY),
            ('experiencias digitales escalables', 17, True, DK),
            (' que conectan tecnología, operaciones y engagement del cliente.', 17, False, LGY),
        ], wrap=True)
        items = [
            ('Mobile & Self-Service Apps', 'Journeys digitales para clientes.'),
            ('Payment & Transactional UX', 'Flujos seguros para conversión.'),
            ('Enterprise Portals', 'Procesos complejos simplificados.'),
            ('B2C Digital Experiences', 'Centrados en usabilidad.'),
            ('B2B Platforms', 'Herramientas operativas.'),
            ('UX/UI Systems', 'Interfaces multi-pantalla.'),
        ]
        footer = ('HUMAN-CENTERED ', 'digital products powered by INFINITE DIGITAL TECHNOLOGY.')
    add_rect(sl, 0.6, 2.3, 0.08, 4.8, OR)
    for idx,(t,d) in enumerate(items):
        y = 2.35 + idx * 0.75
        add_runs(sl, 0.85, y, 6.5, 0.65, [
            (f'• {t}', 14, True, DK),
            (f'  —  {d}', 14, False, LGY),
        ], wrap=True)
    add_runs(sl, 0.6, 6.8, 6.5, 0.5, [
        (footer[0], 13, True, OR),
        (footer[1], 13, False, DK),
    ])
    add_logo(sl, SLIDE_IMGS['como'], 7.1, 0.3, 5.9, 6.9)

def slide_dolores(prs, sl, vertical_nm, pain, lang):
    set_bg(sl, DK)
    add_runs(sl, 0.4, 0.2, 12.0, 0.9, [
        (T(lang,'Dolores en ','Pain Points in '), 32, True, OR),
        (vertical_nm, 32, True, WH),
    ])
    add_rect(sl, 0.4, 1.15, 12.5, 0.05, hc('FF5B0066'))
    add_logo(sl, SLIDE_IMGS['dolores'], 0.4, 1.3, 2.6, 5.9)
    for idx, pt in enumerate(pain):
        col = 0 if idx < 3 else 1
        row = idx % 3
        x = 3.25 + col * 4.8
        y = 1.35 + row * 2.0
        add_rect(sl, x, y, 4.6, 1.85, hc('2E2E2E'), hc('FF5B0055'))
        add_text(sl, x+0.2, y+0.12, 4.2, 0.45,
                 ' '.join(pt.split()[:4]), 13, bold=True, color=OR, wrap=True)
        add_text(sl, x+0.2, y+0.6, 4.2, 1.1, pt, 11, color=MGY, wrap=True)

def slide_case_intro(prs, sl, case, dark, lang):
    bg = DK2 if dark else GR
    tc = WH if dark else DK
    logo_url = LOGO_W if dark else LOGO_B
    set_bg(sl, bg)
    words = case['n'].split()
    mid = (len(words)+1)//2
    line1 = ' '.join(words[:mid])
    line2 = ' '.join(words[mid:]) if len(words) > 1 else ''
    add_text(sl, 0.6, 1.0, 6.5, 1.4, line1, 67, bold=True, color=tc)
    if line2:
        add_text(sl, 0.6, 2.3, 6.5, 1.4, line2, 67, bold=True, color=OR)
    add_text(sl, 0.6, 3.85, 6.3, 1.0, case['sub'], 33, bold=True,
             color=hc('EEEEEE') if dark else DK, wrap=True)
    add_rect(sl, 0.6, 4.95, 1.2, 0.07, OR)
    add_text(sl, 0.6, 5.15, 6.0, 0.35,
             T(lang,'CASO DE ÉXITO · KRUGER WORLDWIDE','SUCCESS STORY · KRUGER WORLDWIDE'),
             13, bold=True, color=hc('999999') if dark else hc('777777'))
    add_logo(sl, logo_url, 0.6, 5.65, 2.0, 0.5)
    if dark:
        add_logo(sl, SLIDE_IMGS['case_intro_a'], 7.1, 0.3, 5.9, 6.9)
    else:
        add_logo(sl, SLIDE_IMGS['case_intro_b'],  7.1, 0.3, 5.9, 6.9)
        add_logo(sl, SLIDE_IMGS['case_intro_b2'], 7.1, 0.3, 2.9, 6.9)

def slide_case_detail(prs, sl, case, lang):
    set_bg(sl, BK)
    add_rect(sl, 0, 0, 13.33, 1.25, GR)
    add_text(sl, 0.4, 0.1, 10.5, 1.0, case['n'], 57, bold=True, color=DK)
    add_logo(sl, LOGO_B, 11.5, 0.3, 1.6, 0.55)
    add_rect(sl, 0, 1.25, 13.33, 0.6, hc('111111'))
    add_text(sl, 0.4, 1.32, 12.0, 0.45, case['sub'].upper(), 16, bold=True, color=WH)
    add_logo(sl, SLIDE_IMGS['case_detail_a'], 0.3, 1.95, 4.8, 5.3)
    if lang == 'en':
        labels = ['CHALLENGE', 'APPROACH', 'OUTCOME']
    else:
        labels = ['RETO', 'ENFOQUE', 'RESULTADO']
    keys = ['ch', 'ap', 'out']
    for idx in range(3):
        y = 2.0 + idx * 1.82
        add_rect(sl, 5.4, y, 7.6, 1.65, hc('1A1A1A'), hc('333333'))
        add_text(sl, 5.65, y+0.15, 4.0, 0.45, labels[idx], 21, bold=True, color=OR2)
        add_text(sl, 5.65, y+0.65, 7.0, 0.9, case[keys[idx]], 13, color=MGY, wrap=True)

def slide_productos(prs, sl, vertical_nm, prods, lang):
    set_bg(sl, GR)
    add_runs(sl, 0.4, 0.1, 11.5, 0.75, [
        (T(lang,'Productos para ','Products for '), 26, True, DK),
        (vertical_nm, 26, True, OR),
    ])
    add_logo(sl, LOGO_B, 11.5, 0.15, 1.6, 0.48)
    for idx, prod in enumerate(prods[:4]):
        x = 0.4 if idx < 2 else 6.85
        y = 1.0 + (idx % 2) * 2.95
        add_rect(sl, x, y, 6.2, 2.75, WH, hc('DDDDDD'))
        add_rect(sl, x, y, 6.2, 0.14, OR)
        add_text(sl, x+0.25, y+0.28, 4.5, 0.5, prod['n'], 18, bold=True, color=DK)
        badge = sl.shapes.add_shape(1, i(x+4.85), i(y+0.28), i(1.1), i(0.38))
        badge.fill.solid(); badge.fill.fore_color.rgb = OR
        badge.line.fill.background()
        add_text(sl, x+4.85, y+0.29, 1.1, 0.36, prod['cat'], 10,
                 bold=True, color=WH, align=PP_ALIGN.CENTER)
        add_text(sl, x+0.25, y+0.85, 5.8, 1.75, prod['d'], 13, color=LGY, wrap=True)

def slide_servicios(prs, sl, vertical_nm, servs, lang):
    set_bg(sl, DK)
    add_runs(sl, 0.4, 0.2, 11.5, 0.85, [
        (T(lang,'Servicios para ','Services for '), 28, True, WH),
        (vertical_nm, 28, True, OR),
    ])
    add_logo(sl, LOGO_W, 11.5, 0.28, 1.6, 0.48)
    add_rect(sl, 0.4, 1.15, 12.5, 0.05, hc('FF5B0066'))
    for idx, s in enumerate(servs):
        x = 0.5 + idx * 3.2
        circ = sl.shapes.add_shape(9, i(x+1.15), i(1.75), i(0.65), i(0.65))
        circ.fill.solid(); circ.fill.fore_color.rgb = OR
        circ.line.fill.background()
        add_text(sl, x+1.15, 1.78, 0.65, 0.6, s['n'], 14,
                 bold=True, color=WH, align=PP_ALIGN.CENTER)
        add_rect(sl, x, 2.55, 3.1, 4.65, hc('2E2E2E'), hc('FF5B0044'))
        add_text(sl, x+0.15, 2.78, 2.8, 0.55, s['t'], 16,
                 bold=True, color=OR, align=PP_ALIGN.CENTER)
        add_text(sl, x+0.15, 3.45, 2.8, 3.5, s['d'], 13,
                 color=MGY, wrap=True, align=PP_ALIGN.CENTER)

def slide_partners(prs, sl, lang):
    set_bg(sl, WH)
    add_logo(sl, LOGO_B, 0.4, 0.2, 2.2, 0.55)
    add_text(sl, 9.0, 0.1, 4.0, 0.85, 'PARTNERS', 42,
             bold=True, color=OR, align=PP_ALIGN.RIGHT)
    # 8 imágenes, una por fila, ancho completo
    img_h = 0.72
    gap = 0.06
    start_y = 1.1
    for idx, url in enumerate(PARTNER_IMGS):
        y = start_y + idx * (img_h + gap)
        try:
            sl.shapes.add_picture(fetch_image(url), i(0.3), i(y), i(12.73), i(img_h))
        except Exception as e:
            print(f"Partner img error {idx}: {e}")
            add_rect(sl, 0.3, y, 12.73, img_h, GR, hc('DDDDDD'))
    add_text(sl, 0, 7.2, 13.33, 0.25,
             T(lang,'14 países · 1500+ proyectos entregados',
                    '14 countries · 1500+ projects delivered'),
             10, color=hc('AAAAAA'), align=PP_ALIGN.CENTER)

def slide_gancho(prs, sl, lang):
    set_bg(sl, WH)
    add_rect(sl, 0, 0, 0.9, 7.5, OR)
    add_text(sl, 1.1, 0.1, 1.0, 1.5, '"', 80, color=hc('FF5B0033'), italic=True)
    if lang == 'en':
        add_text(sl, 1.2, 1.5, 7.2, 3.0,
                 '"The future of business is digital. The difference between leading and falling behind lies in the experience you offer your customers."',
                 17, bold=True, italic=True, color=DK, wrap=True)
        add_text(sl, 1.2, 4.85, 7.5, 0.9,
                 'Kruger has transformed more than 1500 digital products across 14 countries.',
                 13, color=LGY, wrap=True)
    else:
        add_text(sl, 1.2, 1.5, 7.2, 3.0,
                 '"El futuro de los negocios es digital. La diferencia entre liderar y quedarse atrás está en la experiencia que ofreces a tus clientes."',
                 17, bold=True, italic=True, color=DK, wrap=True)
        add_text(sl, 1.2, 4.85, 7.5, 0.9,
                 'Kruger ha transformado más de 1500 productos digitales en 14 países.',
                 13, color=LGY, wrap=True)
    add_rect(sl, 1.2, 4.65, 1.5, 0.08, OR)
    add_logo(sl, SLIDE_IMGS['gancho_main'], 9.1, 0.3, 3.9, 6.9)
    add_logo(sl, SLIDE_IMGS['gancho_side'], 9.1, 0.3, 1.9, 6.9)

def slide_cierre(prs, sl, cfg, lang):
    set_bg(sl, OR)
    add_text(sl, 0.4, 0.28, 7.0, 0.42, cfg.get('name',''), 13, bold=True, color=WH)
    add_text(sl, 0.4, 0.7, 7.0, 0.35, cfg.get('title',''), 11, color=hc('FFD0B0'))
    add_text(sl, 7.0, 0.28, 5.9, 0.42, 'KRUGER CORPORATION', 13,
             bold=True, color=WH, align=PP_ALIGN.RIGHT)
    add_text(sl, 7.0, 0.7, 5.9, 0.35, 'krugerworldwide.com', 11,
             color=hc('FFD0B0'), align=PP_ALIGN.RIGHT)
    add_logo(sl, LOGO_W, 4.8, 1.5, 3.7, 0.92)
    add_rect(sl, 3.0, 2.65, 7.3, 0.06, hc('FFFFFF88'))
    add_logo(sl, SLIDE_IMGS['cierre'], 5.65, 2.85, 2.0, 2.0)
    email = cfg.get('email','worldwide@krugercorp.com')
    phone = cfg.get('phone','')
    contact = f"{email}  ·  {phone}" if phone else email
    add_text(sl, 0, 5.1, 13.33, 0.45, contact, 15, bold=True, color=WH, align=PP_ALIGN.CENTER)
    if cfg.get('client'):
        add_text(sl, 0, 5.65, 13.33, 0.38,
                 f"Preparado especialmente para {cfg['client']}" if lang=='es' else f"Specially prepared for {cfg['client']}",
                 12, color=hc('FFE8D6'), align=PP_ALIGN.CENTER, italic=True)
    add_text(sl, 0.4, 7.2, 5.0, 0.22, '© 2026 KRUGER CORP.', 8, color=WH)
    add_text(sl, 8.0, 7.2, 5.0, 0.22,
             T(lang,'Confidencial','Confidential'), 8, color=WH, align=PP_ALIGN.RIGHT)

# ── ENDPOINT ─────────────────────────────────────────────────
@app.route('/generate', methods=['POST','OPTIONS'])
def generate():
    if request.method == 'OPTIONS':
        return '', 200
    data = request.json
    vertical = data.get('vertical', {})
    cfg = data.get('cfg', {})
    lang = data.get('lang', 'es')  # 'es' o 'en'

    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    def ns(): return prs.slides.add_slide(blank)

    slide_portada(prs, ns(), cfg, lang)
    slide_disclaimer(prs, ns(), lang)
    slide_quienes(prs, ns(), lang)
    slide_porque(prs, ns(), lang)
    slide_como(prs, ns(), lang)
    slide_dolores(prs, ns(), vertical.get('nm',''), vertical.get('pain',[]), lang)
    for idx, case in enumerate(vertical.get('cases', [])):
        slide_case_intro(prs, ns(), case, idx % 2 == 0, lang)
        slide_case_detail(prs, ns(), case, lang)
    slide_productos(prs, ns(), vertical.get('nm',''), vertical.get('prods',[]), lang)
    slide_servicios(prs, ns(), vertical.get('nm',''), vertical.get('servs',[]), lang)
    slide_partners(prs, ns(), lang)
    slide_gancho(prs, ns(), lang)
    slide_cierre(prs, ns(), cfg, lang)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    nm = vertical.get('nm','Vertical').replace(' ','_')
    cl = cfg.get('client','Cliente').replace(' ','_')
    lang_suffix = 'EN' if lang == 'en' else 'ES'
    filename = f'Kruger_{nm}_{cl}_2026_{lang_suffix}.pptx'

    return send_file(buf,
        mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
        as_attachment=True,
        download_name=filename)

@app.route('/health', methods=['GET'])
def health():
    return jsonify({'status': 'ok'})

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=False)
